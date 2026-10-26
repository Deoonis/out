import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

from jobassist import ai, jobs
from jobassist.store import Store

HH_ITEM = {"id": "123", "name": "Менеджер проектов", "employer": {"name": "Компания А"}, "area": {"name": "Москва"},
           "salary": {"from": 200000, "to": 250000, "currency": "RUR"}, "schedule": {"name": "Удаленная работа"},
           "alternate_url": "https://hh.ru/vacancy/123", "published_at": "2026-10-20T10:00:00+0300",
           "snippet": {"requirement": "Опыт <highlighttext>Agile</highlighttext>, управление бюджетом",
                       "responsibility": "Ведение B2B-проектов"}}

PROFILE = {"roles": "менеджер проектов, project manager", "keywords": "agile, бюджет, b2b",
           "stopwords": "стажёр, холодные звонки", "salary": "220 000", "formats": ["удалёнка"]}


class JobsTest(unittest.TestCase):
    def test_normalize(self):
        v = jobs.normalize_hh(HH_ITEM)
        self.assertEqual((v["id"], v["published"]), ("hh-123", "2026-10-20"))
        self.assertIn("Agile", v["text"])
        self.assertNotIn("<", v["text"])

    def test_score(self):
        good = jobs.score_vacancy(jobs.normalize_hh(HH_ITEM), PROFILE)
        bad = jobs.score_vacancy({"title": "Стажёр", "text": "холодные звонки", "salary_to": 40000}, PROFILE)
        self.assertGreaterEqual(good["score"], 85, good)
        self.assertEqual(bad["score"], 0, bad)

    def test_score_disabled_without_profile(self):
        self.assertIsNone(jobs.score_vacancy(jobs.normalize_hh(HH_ITEM), {})["score"])

    def test_export_xlsx(self):
        data = jobs.export_xlsx({"vacancies": [{"company": "А", "title": "PM", "stage": "3"}],
                                 "skills": [{"block": "Планирование", "name": "Оценка задач", "level": 6}]})
        with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
            f.write(data)
        wb = openpyxl.load_workbook(f.name)
        os.unlink(f.name)
        self.assertEqual(wb.sheetnames, ["Отклики", "Навыки", "Отказы"])
        self.assertEqual(wb["Отклики"]["E2"].value, "Интервью с HR")


class StoreTest(unittest.TestCase):
    def test_roundtrip(self):
        with tempfile.TemporaryDirectory() as d:
            st = Store(Path(d) / "t.db")
            self.assertEqual(st.load(), {})
            st.save({"a": [1, "б"], "b": {"x": 1}})
            st.save({"a": [2]})
            self.assertEqual(st.load(), {"a": [2]})


class AiTest(unittest.TestCase):
    def test_key_env_wins_over_settings(self):
        with mock.patch.dict(os.environ, {"GEMINI_API_KEY": "env-key"}):
            self.assertEqual(ai.resolve_key({"ai_provider": "gemini", "ai_key": "field"}), "env-key")
        with mock.patch.dict(os.environ, {}, clear=True):
            self.assertEqual(ai.resolve_key({"ai_provider": "gemini", "ai_key": " field "}), "field")
            self.assertEqual(ai.resolve_key({"ai_provider": "ollama"}), "local")
            with self.assertRaises(RuntimeError):
                ai.resolve_key({"ai_provider": "openai"})

    def test_chat_request(self):
        seen = {}

        def fake(url, body=None, headers=None, timeout=30):
            seen.update(url=url, body=body, headers=headers)
            return {"choices": [{"message": {"content": " Письмо "}}]}

        with mock.patch.object(ai, "http_json", fake), mock.patch.dict(os.environ, {}, clear=True):
            out = ai.run_task("cover_letter", "вакансия", "резюме", {"title": "PM"},
                              {"ai_provider": "gemini", "ai_key": "k"})
        self.assertEqual(out, "Письмо")
        self.assertTrue(seen["url"].endswith("/v1beta/openai/chat/completions"))
        self.assertEqual(seen["headers"]["Authorization"], "Bearer k")
        self.assertEqual(seen["body"]["model"], "gemini-2.5-flash")
        self.assertIn("вакансия", json.dumps(seen["body"], ensure_ascii=False))


if __name__ == "__main__":
    unittest.main()


class ResumeTest(unittest.TestCase):
    def setUp(self):
        from jobassist import resume
        self.resume = resume
        self.text = (Path(__file__).parent / "fixture_hh_resume.txt").read_text(encoding="utf-8")

    def test_parse_hh(self):
        r = self.resume.parse_hh(self.text)
        self.assertTrue(r["is_hh"])
        self.assertEqual(r["title"], "Руководитель проектов")
        self.assertEqual(r["specializations"], ["Руководитель проектов", "Менеджер продукта"])
        self.assertEqual(r["salary"], "250000")
        self.assertEqual(r["formats"], ["гибрид", "удалёнка"])
        self.assertEqual(r["city"], "Казань")
        self.assertEqual(r["experience_total"], "8 лет 2 месяца")
        self.assertEqual(r["skills"], ["Agile", "Scrum", "MS Project", "Управление рисками", "Бюджетирование", "Jira"])
        self.assertEqual([(j["company"], j["title"]) for j in r["jobs"]],
                         [("ООО Пример", "Руководитель проектов"), ("Тестовая Компания", "Менеджер проектов")])
        self.assertEqual(r["jobs"][0]["period"], "Март 2022 — по настоящее время")
        self.assertTrue(r["about"].startswith("Руководитель проектов с опытом"))
        self.assertNotIn("Резюме обновлено", " ".join(self.resume.clean_lines(self.text)))

    def test_non_hh_and_formats(self):
        self.assertFalse(self.resume.parse_hh("Просто текст резюме")["is_hh"])
        self.assertEqual(self.resume.extract_text("a.txt", "привет".encode("cp1251")), "привет")
        with self.assertRaises(ValueError):
            self.resume.extract_text("a.rtf", b"{}")

    def test_docx(self):
        import docx
        d = docx.Document()
        for line in self.text.splitlines():
            d.add_paragraph(line)
        buf = __import__("io").BytesIO()
        d.save(buf)
        self.assertEqual(self.resume.parse_hh(self.resume.extract_text("cv.docx", buf.getvalue()))["title"],
                         "Руководитель проектов")


class AiJsonTest(unittest.TestCase):
    def test_parse_json_fenced(self):
        self.assertEqual(ai.parse_json('Вот:\n```json\n{"a": [1]}\n```'), {"a": [1]})
        self.assertEqual(ai.parse_json('ответ {"requirements": ["Agile"]}'), {"requirements": ["Agile"]})
        with self.assertRaises(RuntimeError):
            ai.parse_json("нет json")

    def test_grill_uses_methodology(self):
        seen = {}
        with mock.patch.object(ai, "chat", lambda msgs, s: seen.setdefault("m", msgs) and "ok"):
            ai.grill("резюме", "цель", ["в1", "в2"], {})
        self.assertIn("XYZ", seen["m"][0]["content"])
        self.assertIn("(3) ВАКАНСИИ", seen["m"][1]["content"])
